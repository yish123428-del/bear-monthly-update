"""현황판 갱신 스크립트 테스트 (합성 베이스 사용)."""
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / ".claude" / "skills" / "bear-monthly-update" / "scripts"))
sys.path.insert(0, str(ROOT / "tests"))

try:
    import openpyxl
    HAVE_OPENPYXL = True
except ImportError:
    HAVE_OPENPYXL = False


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl 미설치")
class TestUpdate(unittest.TestCase):
    def setUp(self):
        import helpers
        import update_status_excel as use
        self.use = use
        self.tmp = tempfile.TemporaryDirectory()
        self.base = helpers.make_base_xlsx(Path(self.tmp.name) / "base.xlsx",
                                           counts={"소화기2사업팀": {"4월": 1}})
        self.out = Path(self.tmp.name) / "out.xlsx"
        self.proposals = [
            {"월": "26년 8월", "본부": "MKT", "사업부": "마케팅1사업부", "팀": "소화기2사업팀", "제안자": "김혁(C)",
             "건수": 1, "제안유형": "의약품 자체개발", "아이디어유형": "복합제", "제안내용": "NSAIDs + 소화기용제 복합제",
             "검토결과": "1차 타당성 검토중", "이메일표기": ""},
            {"월": "26년 8월", "본부": "영업본부", "사업부": "서울3사업부", "팀": "북부2", "제안자": "최태우",
             "건수": 1, "제안유형": "의약품 도입/제휴", "아이디어유형": "개량신약/제네릭", "제안내용": "엔커버",
             "검토결과": "검토 중단", "이메일표기": ""},
            {"월": "26년 8월", "본부": "MKT", "사업부": "마케팅2사업부", "팀": "미지의팀사업팀", "제안자": "홍길동",
             "건수": 1, "제안유형": "의료기기 도입/제휴", "아이디어유형": "기타", "제안내용": "제품X",
             "검토결과": "유관부서 검토중", "이메일표기": ""},
        ]
        self.warnings = use.run(str(self.base), self.proposals, "26년 9월", "8월", 3, 87, str(self.out))
        self.wb = openpyxl.load_workbook(self.out)
        self.ws1 = self.wb["1. 제안 현황"]
        self.ws2 = self.wb["2. 누적 제안 요약(26.01~)"]

    def tearDown(self):
        self.tmp.cleanup()

    def find_row(self, name):
        for r in range(1, self.ws1.max_row + 1):
            if self.ws1.cell(row=r, column=3).value == name:
                return r
        raise AssertionError(name)

    def test_title_and_total_text(self):
        self.assertEqual(self.ws1["A1"].value, "사업팀 별 아이디어 현황판 (26년 9월)")
        self.assertEqual(self.ws1["A3"].value, "[ 누계 진행 결과 _8월 3건 총 87건 ] ")
        self.assertEqual(self.ws2["A1"].value, "누적 제안 요약 (26년 9월)")

    def test_month_count_incremented(self):
        r = self.find_row("소화기2사업팀")
        self.assertEqual(self.ws1.cell(row=r, column=6 + 7).value, 1)   # 8월 = M열
        self.assertEqual(self.ws1.cell(row=r, column=6 + 3).value, 1)   # 기존 4월 값 유지
        r2 = self.find_row("북부2")
        self.assertEqual(self.ws1.cell(row=r2, column=6 + 7).value, 1)
        # 영업부 제안은 사업부별 표도 +1, MKT 제안의 사업부(마케팅1사업부)는 표가 없으므로 경고 없이 무시
        r3 = self.find_row("서울3사업부")
        self.assertEqual(self.ws1.cell(row=r3, column=6 + 7).value, 1)
        self.assertEqual(self.ws1.cell(row=self.find_row("서울2사업부"), column=6 + 7).value, 0)

    def test_hidden_legacy_header_skipped(self):
        import helpers
        import update_status_excel as use
        base = helpers.make_base_xlsx(Path(self.tmp.name) / "b2.xlsx")
        wb = openpyxl.load_workbook(base)
        ws = wb["1. 제안 현황"]
        # 상단에 숨김 처리된 구버전 표를 흉내 낸다 (row 5: 헤더 '팀', row 6: 데이터)
        ws.cell(row=5, column=3, value="팀")
        ws.cell(row=5, column=6, value="1월")
        for i in range(12):
            ws.cell(row=5, column=6 + i, value=f"{i + 1}월")
        ws.cell(row=6, column=3, value="소화기2사업팀")
        ws.cell(row=6, column=13, value=7)
        ws.row_dimensions[5].hidden = True
        ws.row_dimensions[6].hidden = True
        wb.save(base)
        out = Path(self.tmp.name) / "o2.xlsx"
        use.run(str(base), [self.proposals[0]], "26년 9월", "8월", 1, 85, str(out))
        ws = openpyxl.load_workbook(out)["1. 제안 현황"]
        self.assertEqual(ws.cell(row=6, column=13).value, 7)   # 숨김 표는 그대로
        r = next(r for r in range(10, ws.max_row + 1) if ws.cell(row=r, column=3).value == "소화기2사업팀")
        self.assertEqual(ws.cell(row=r, column=13).value, 1)

    def test_formulas_preserved(self):
        r = self.find_row("소화기2사업팀")
        self.assertTrue(str(self.ws1.cell(row=r, column=4).value).startswith("=SUM("))
        self.assertTrue(str(self.ws1.cell(row=r, column=5).value).startswith("=SUM("))

    def test_unmatched_warning(self):
        self.assertTrue(any("매칭 실패" in w and "미지의팀사업팀" in w for w in self.warnings))

    def test_sheet2_rows_appended(self):
        rows = [self.ws2.cell(row=r, column=1).value for r in range(4, self.ws2.max_row + 1)]
        self.assertEqual(len(rows), 2 + 3)
        last = [self.ws2.cell(row=self.ws2.max_row, column=c).value for c in range(1, 8)]
        self.assertEqual(last, ["26년 8월", "미지의팀사업팀", "홍길동", "의료기기 도입/제휴", "기타", "제품X", "유관부서 검토중"])


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl 미설치")
class TestHtmlRender(unittest.TestCase):
    def test_formula_evaluation_and_html(self):
        import helpers
        import xlsx_html
        with tempfile.TemporaryDirectory() as tmp:
            base = helpers.make_base_xlsx(Path(tmp) / "b.xlsx", counts={"소화기2사업팀": {"4월": 1, "8월": 2}})
            wb = openpyxl.load_workbook(base)
            ws = wb["1. 제안 현황"]
            ev = xlsx_html.FormulaEvaluator(ws)
            r = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=3).value == "소화기2사업팀")
            self.assertEqual(ev.evaluate(f"SUM(F{r}:Q{r})"), 3)
            self.assertEqual(ev.evaluate(f"I{r}+M{r}"), 3)   # 4월(I) 1 + 8월(M) 2
            self.assertEqual(ev.evaluate(f"F{r}+M{r}"), 2)
            doc, warnings = xlsx_html.sheet_to_html(ws, Path(tmp), stop_regex=r"사무소별|^사무소$")
            self.assertIn("소화기2사업팀", doc)
            self.assertNotIn("병원경기2", doc)   # 3. 사무소별 섹션은 제외
            self.assertIn("2. 영업부 사업부별", doc)
            self.assertEqual(warnings, [])


if __name__ == "__main__":
    unittest.main()
