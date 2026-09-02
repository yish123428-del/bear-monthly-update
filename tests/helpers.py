"""테스트용 합성 베이스 현황판 생성.

실제 현황판(BEAR_아이디어_제안_현황판_YY년M월_v1.xlsx)의 구조를 update_status_excel.py 가
기대하는 최소 형태로 흉내 낸다.

시트 '1. 제안 현황'
  A1  제목 "사업팀 별 아이디어 현황판 (26년 8월)"
  A3  "[ 누계 진행 결과 _7월 4건 총 84건 ]"
  섹션마다 2행 헤더: (r)   팀 | 제안 합계 | 누계(26년) | 26년 ...
                     (r+1)      |          |            | 1월 | 2월 | ... | 12월
                     (r+2) 빈 줄, (r+3)~ 데이터  ← 스크립트가 서브헤더를 보면 header+3 부터 읽는다
시트 '2. 누적 제안 요약(26.01~)'
  A1 제목 "누적 제안 요약 (26년 8월)", 3행 헤더, 4행~ 데이터
"""
from __future__ import annotations

from pathlib import Path

MKT_TEAMS = ["소화기1사업팀", "소화기2사업팀", "순환기1사업팀", "순환기2사업팀",
             "호흡피부안과사업팀", "디지털헬스1사업팀", "신경골격사업팀"]
SALES_DIVISIONS = ["서울1", "서울2", "서울3", "서울4", "지방2"]
SALES_OFFICES = ["병원경기2", "북부2", "병원광주2", "경인3"]

SHEET2_HEADERS = ["제안 연월", "사업팀/사무소", "제안자", "제안 유형", "아이디어 유형", "제안 내용", "검토 결과"]


def _section(ws, start_row: int, label: str, names: list[str], counts: dict[str, dict[str, int]] | None = None):
    """한 섹션을 그리고 다음 섹션의 시작 행을 돌려준다."""
    counts = counts or {}
    ws.cell(row=start_row, column=3, value=label)
    ws.cell(row=start_row, column=4, value="제안 합계")
    ws.cell(row=start_row, column=5, value="누계(26년)")
    ws.cell(row=start_row, column=6, value="26년")
    for i in range(12):
        ws.cell(row=start_row + 1, column=6 + i, value=f"{i + 1}월")
    r = start_row + 3
    for name in names:
        ws.cell(row=r, column=3, value=name)
        ws.cell(row=r, column=4, value=f"=SUM(F{r}:Q{r})")
        ws.cell(row=r, column=5, value=f"=SUM(F{r}:Q{r})")
        for i in range(12):
            v = counts.get(name, {}).get(f"{i + 1}월")
            ws.cell(row=r, column=6 + i, value=v if v is not None else 0)
        r += 1
    return r + 2


def make_base_xlsx(path: str | Path, report_month="26년 8월", prev_text="[ 누계 진행 결과 _7월 4건 총 84건 ] ",
                   sheet2_rows=None, counts=None) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "1. 제안 현황"
    ws1["A1"] = f"사업팀 별 아이디어 현황판 ({report_month})"
    ws1["A3"] = prev_text
    r = 10
    r = _section(ws1, r, "팀", MKT_TEAMS, counts)
    ws1.cell(row=r - 1, column=1, value="2. 영업부 사업부별")
    r = _section(ws1, r, "사업부", SALES_DIVISIONS, counts)
    ws1.cell(row=r - 1, column=1, value="3. 영업부 사무소별")
    _section(ws1, r, "사무소", SALES_OFFICES, counts)

    ws2 = wb.create_sheet("2. 누적 제안 요약(26.01~)")
    ws2["A1"] = f"누적 제안 요약 ({report_month})"
    for c, h in enumerate(SHEET2_HEADERS, 1):
        ws2.cell(row=3, column=c, value=h)
    rows = sheet2_rows if sheet2_rows is not None else [
        ["26년 4월", "소화기2사업팀", "김민수", "의약품 자체개발", "복합제", "우루사+실리스칸 복합제(UDCA+실리마린)", "1차 타당성 검토중"],
        ["26년 7월", "병원경기2", "권봉기", "의약품 자체개발", "복합제", "펙수클루+클로아트+리토바젯 복합제", "1차 타당성 검토중"],
    ]
    for r_i, row in enumerate(rows, 4):
        for c, v in enumerate(row, 1):
            ws2.cell(row=r_i, column=c, value=v)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
