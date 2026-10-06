"""시트 덤프 → 신규 제안 변환 테스트."""
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SKILL = ROOT / ".claude" / "skills" / "bear-monthly-update"
sys.path.insert(0, str(SKILL / "scripts"))
sys.path.insert(0, str(ROOT / "tests"))

try:
    import openpyxl  # noqa: F401
    HAVE_OPENPYXL = True
except ImportError:
    HAVE_OPENPYXL = False

FIXTURE = ROOT / "tests" / "fixtures" / "sheet_dump_sample.md"
CONFIG = json.loads((SKILL / "config.json").read_text(encoding="utf-8"))


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl 미설치")
class TestParseDump(unittest.TestCase):
    def setUp(self):
        import dumpio
        import sheet_to_proposals as s2p
        self.s2p = s2p
        self.rows, self.warnings = s2p.parse_dump(dumpio.load_dump(FIXTURE), CONFIG["sheet"]["anchor_labels"])

    def test_first_block_only(self):
        months = [r["제안월"] for r in self.rows]
        self.assertNotIn("25년 1월", months, "두 번째(구버전) 탭이 섞이면 안 됨")
        self.assertEqual(len(self.rows), 8)

    def test_month_normalized(self):
        self.assertEqual(self.rows[0]["제안월"], "24년 11월")   # '24냔 11월' 오타
        self.assertEqual(self.rows[3]["제안월"], "26년 7월")    # '26년  7월' 이중 공백

    def test_columns_aligned(self):
        r = self.rows[1]
        self.assertEqual(r["접수자"], "김민수")
        self.assertEqual(r["KOL"], "김민수PM")
        self.assertEqual(r["대분류"], "복합제")
        self.assertEqual(r["현황"], "진행중")
        self.assertEqual(r["세부현황"], "1차 타당성 검토")
        self.assertEqual(self.rows[0]["소분류"], "-")  # \- 언이스케이프

    def test_csv_dump_matches_markdown(self):
        """download_file_content(text/csv) 경로가 markdown 경로와 같은 레코드를 낸다."""
        import csv
        import io
        import json
        import base64
        import dumpio
        text = dumpio.load_dump(FIXTURE)
        lines = text.split("\n")
        hdr_idx, header = dumpio.find_header_blocks(lines, CONFIG["sheet"]["anchor_labels"])[0]
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(["■ Bear 아이디어 제안(베아제) 현황판"] + [""] * 21)
        w.writerow([""] * 22)
        # 셀 안 줄바꿈이 있는 헤더도 처리되는지: '사업팀(MKT)\n 사무소(영업)'
        w.writerow([h.replace("사업팀(MKT) 사무소(영업)", "사업팀(MKT)\n 사무소(영업)") for h in header])
        for line in lines[hdr_idx + 1:]:
            if "|" not in line:
                break
            cells = dumpio.split_row(line)
            if dumpio.is_separator(cells):
                continue
            w.writerow([dumpio.unescape(c) for c in cells])
        csv_text = buf.getvalue()
        self.assertTrue(dumpio.is_csv_dump(csv_text))
        rows_csv, _ = self.s2p.parse_dump(csv_text, CONFIG["sheet"]["anchor_labels"])
        self.assertEqual(len(rows_csv), len(self.rows))
        self.assertEqual(rows_csv[1], self.rows[1])
        self.assertEqual(rows_csv[0]["제안월"], "24년 11월")
        # download_file_content JSON 포장(base64) 도 load_dump 가 푼다
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "dl.json"
            p.write_text(json.dumps({"content": base64.b64encode(csv_text.encode("utf-8")).decode(),
                                     "mimeType": "text/csv", "title": "x", "id": "y"}), encoding="utf-8")
            self.assertEqual(dumpio.load_dump(p), csv_text)


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl 미설치")
class TestNormalize(unittest.TestCase):
    def setUp(self):
        import sheet_to_proposals as s2p
        self.s2p = s2p
        self.m = CONFIG["mapping"]

    def test_bonbu(self):
        for raw, want in [("ETC마케팅본부", "MKT"), ("ETC 마케팅", "MKT"), ("ETC병원", "영업본부"),
                          ("ETC 병원본부", "영업본부"), ("ETC로컬본부", "영업본부")]:
            self.assertEqual(self.s2p.norm_bonbu(raw, self.m), want)

    def test_team_alias_and_suffix(self):
        self.assertEqual(self.s2p.norm_team("호피안", "MKT", self.m), "호흡피부안과사업팀")
        self.assertEqual(self.s2p.norm_team("호흡피부안과", "MKT", self.m), "호흡피부안과사업팀")
        self.assertEqual(self.s2p.norm_team("소화기2", "MKT", self.m), "소화기2사업팀")
        self.assertEqual(self.s2p.norm_team("병원경기2", "영업본부", self.m), "병원경기2")
        self.assertEqual(self.s2p.norm_division("마케팅1", "MKT", self.m), "마케팅1사업부")
        self.assertEqual(self.s2p.norm_division("서울2", "영업본부", self.m), "서울2사업부")
        self.assertEqual(self.s2p.norm_team("디지털1", "MKT", self.m), "디지털헬스1팀")

    def test_display_fields_for_sales(self):
        row = {"제안월": "26년 7월", "본부": "ETC 병원본부", "사업부": "서울2", "팀": "병원경기2", "접수자": "권봉기",
               "품목분류": "의약품", "개발형태": "자체 개발", "대분류": "복합제", "현황": "진행중", "세부현황": "1차 타당성 검토"}
        p = self.s2p.to_proposal(row, self.m)
        self.assertEqual((p["본부"], p["본부표시"]), ("영업본부", ""))
        self.assertEqual((p["팀"], p["팀표시"]), ("병원경기2", "병원경기2사무소"))
        self.assertEqual(p["사업부"], "서울2사업부")
        mkt = self.s2p.to_proposal({**row, "본부": "ETC마케팅본부", "사업부": "마케팅1", "팀": "소화기2"}, self.m)
        self.assertEqual((mkt["본부표시"], mkt["팀표시"]), ("MKT", "소화기2사업팀"))

    def test_review_status(self):
        self.assertEqual(self.s2p.norm_review("1차 타당성 검토", "진행중", self.m), "1차 타당성 검토중")
        self.assertEqual(self.s2p.norm_review("1차 피드백", "검토 중단", self.m), "검토 중단")
        self.assertEqual(self.s2p.norm_review("NBTS 제안", "진행중", self.m), "NBTS 제안")

    def test_proposal_fields(self):
        row = {"제안월": "26년 4월", "본부": "ETC마케팅", "사업부": "마케팅1", "팀": "소화기2", "접수자": "김민수",
               "품목분류": "의약품", "개발형태": "자체 개발", "대분류": "복합제", "소분류": "복용편의성 개선",
               "제품명": "우루사 100mg 실리스칸 175mg", "성분": "UDCA+실리마린", "현황": "진행중",
               "세부현황": "1차 타당성 검토"}
        p = self.s2p.to_proposal(row, self.m)
        self.assertEqual(p["제안유형"], "의약품 자체개발")
        self.assertEqual(p["아이디어유형"], "복합제")
        self.assertEqual(p["제안자"], "김민수")
        self.assertEqual(p["건수"], 1)
        self.assertEqual(p["제안내용"], "우루사 100mg 실리스칸 175mg 복합제(UDCA+실리마린)")
        self.assertTrue(p["_draft"])


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl 미설치")
class TestSelectionEndToEnd(unittest.TestCase):
    def setUp(self):
        import helpers
        import sheet_to_proposals as s2p
        self.s2p = s2p
        self.tmp = tempfile.TemporaryDirectory()
        self.base = helpers.make_base_xlsx(Path(self.tmp.name) / "BEAR_아이디어_제안_현황판_26년8월_v1.xlsx")
        self.meta = s2p.run(str(FIXTURE), "26년 9월", Path(self.tmp.name) / "work", CONFIG, str(self.base))
        self.data = json.loads((Path(self.tmp.name) / "work" / "proposals.json").read_text(encoding="utf-8"))
        self.props = self.data["proposals"]

    def tearDown(self):
        self.tmp.cleanup()

    def keys(self):
        return [(p["월"], p["제안자"], p["_reason"]) for p in self.props]

    def test_selected_rows(self):
        keys = self.keys()
        self.assertIn(("26년 8월", "김혁(C)", "전월 신규"), keys)
        self.assertIn(("26년 8월", "최태우", "전월 신규"), keys)
        self.assertIn(("26년 8월", "홍길동", "전월 신규"), keys)
        # 26년 4월 이영민은 베이스 시트2에 없어 캐치업, 김민수는 이미 있어 제외
        self.assertTrue(any(k[1] == "이영민" and "캐치업" in k[2] for k in keys))
        self.assertFalse(any(k[1] == "김민수" for k in keys))
        self.assertFalse(any(k[1] == "권봉기" for k in keys))
        self.assertEqual(len(self.props), 4)

    def test_dedupe(self):
        self.assertEqual(sum(1 for p in self.props if p["제안자"] == "김혁(C)"), 1)
        self.assertTrue(any("중복 행 제거" in w for w in self.meta["warnings"]))

    def test_unmatched_flag(self):
        by = {p["제안자"]: p for p in self.props}
        self.assertTrue(by["홍길동"]["_unmatched"])
        self.assertEqual(by["홍길동"]["팀"], "미지의팀사업팀")
        self.assertFalse(by["김혁(C)"]["_unmatched"])
        self.assertFalse(by["최태우"]["_unmatched"])  # 북부2 는 사무소 섹션에 있음
        self.assertTrue(any("팀 미매칭" in w for w in self.meta["warnings"]))

    def test_counts_and_warnings(self):
        self.assertEqual(self.meta["new_count"], 4)
        self.assertEqual(self.meta["total_count"], 84 + 4)
        self.assertEqual(self.meta["sheet_row_count"], 8)
        self.assertTrue(any("누계 불일치" in w for w in self.meta["warnings"]))

    def test_long_ingredient_truncated(self):
        by = {p["제안자"]: p for p in self.props}
        self.assertLess(len(by["최태우"]["제안내용"]), 120)
        self.assertIn("Casein", by["최태우"]["_source"]["성분"])  # 원문은 보존

    def test_outputs_exist(self):
        work = Path(self.tmp.name) / "work"
        self.assertTrue((work / "new_proposals.xlsx").exists())
        self.assertTrue((work / "review.md").exists())
        review = (work / "review.md").read_text(encoding="utf-8")
        self.assertIn("팀 미매칭", review)
        wb = openpyxl.load_workbook(work / "new_proposals.xlsx")
        ws = wb.active
        self.assertEqual(ws.cell(row=3, column=1).value, "제안 연월")
        self.assertEqual(ws.max_row, 3 + 4)

    def test_no_base(self):
        meta = self.s2p.run(str(FIXTURE), "26년 9월", Path(self.tmp.name) / "work2", CONFIG, None)
        self.assertEqual(meta["new_count"], 3)          # 캐치업 없음
        self.assertEqual(meta["total_count"], 8)        # 시트 행 수 사용
        self.assertEqual(meta["total_count_source"], "sheet_rows")


if __name__ == "__main__":
    unittest.main()


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl 미설치")
class TestStatusSync(unittest.TestCase):
    """시트에서 '검토 중단' 으로 바뀐 기존 건을 누적 요약(시트2) 검토 결과에 반영."""

    def setUp(self):
        import helpers
        import sheet_to_proposals as s2p
        import update_status_excel as use
        self.s2p, self.use = s2p, use
        self.tmp = tempfile.TemporaryDirectory()
        rows = [
            # 시트: 진행중 → 그대로 둬야 함 (사람이 다듬은 '유관부서 검토중' 유지)
            ["26년 4월", "소화기2사업팀", "김민수", "의약품 자체개발", "복합제", "우루사+실리스칸 복합제(UDCA+실리마린)", "유관부서 검토중"],
            # 시트: 검토 중단 → 갱신 대상
            ["26년 8월", "북부2", "최태우", "의약품 도입/제휴", "개량신약/제네릭", "엔커버 제네릭", "1차 타당성 검토중"],
            # 이미 검토 중단 → 변경 없음
            ["26년 7월", "병원경기2", "권봉기", "의약품 자체개발", "복합제", "펙수클루+클로아트+리토바젯 복합제", "검토 중단"],
        ]
        self.base = helpers.make_base_xlsx(Path(self.tmp.name) / "BEAR_아이디어_제안_현황판_26년9월_v1.xlsx",
                                           report_month="26년 9월", sheet2_rows=rows)

    def tearDown(self):
        self.tmp.cleanup()

    def test_only_stopped_rows_updated(self):
        meta = self.s2p.run(str(FIXTURE), "26년 10월", Path(self.tmp.name) / "w", CONFIG, str(self.base))
        ups = meta["status_updates"]
        self.assertEqual(len(ups), 1, ups)
        self.assertEqual(ups[0]["old"], "1차 타당성 검토중")
        self.assertEqual(ups[0]["new"], "검토 중단")
        self.assertIn("최태우", ups[0]["label"])
        review = (Path(self.tmp.name) / "w" / "review.md").read_text(encoding="utf-8")
        self.assertIn("누적 요약 검토 결과 갱신", review)

        out = Path(self.tmp.name) / "out.xlsx"
        warns = self.use.run(str(self.base), [], "26년 10월", "9월", 0, 85, str(out), status_updates=ups)
        self.assertEqual(warns, [])
        ws2 = openpyxl.load_workbook(out)["2. 누적 제안 요약(26.01~)"]
        vals = {ws2.cell(row=r, column=3).value: ws2.cell(row=r, column=7).value for r in range(4, 7)}
        self.assertEqual(vals["최태우"], "검토 중단")
        self.assertEqual(vals["김민수"], "유관부서 검토중")
        self.assertEqual(ws2["A1"].value, "누적 제안 요약 (26년 10월)")
        ws1 = openpyxl.load_workbook(out)["1. 제안 현황"]
        self.assertEqual(ws1["A3"].value, "[ 누계 진행 결과 _9월 0건 총 85건 ] ")

    def test_stale_old_value_is_skipped(self):
        out = Path(self.tmp.name) / "out2.xlsx"
        bad = [{"row": 5, "old": "다른 값", "new": "검토 중단", "label": "x"}]
        warns = self.use.run(str(self.base), [], "26년 10월", "9월", 0, 85, str(out), status_updates=bad)
        self.assertTrue(any("불일치" in w for w in warns))
        ws2 = openpyxl.load_workbook(out)["2. 누적 제안 요약(26.01~)"]
        self.assertEqual(ws2.cell(row=5, column=7).value, "1차 타당성 검토중")
