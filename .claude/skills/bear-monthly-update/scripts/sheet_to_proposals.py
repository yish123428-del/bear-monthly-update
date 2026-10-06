#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Google Sheet 「Bear 아이디어 제안」 덤프 → 스킬 입력(신규 제안) 변환.

입력  : mcp__Google_Drive__read_file_content 결과 (JSON 또는 markdown), 이전 달 현황판(.xlsx)
출력  : new_proposals.xlsx (references/input_format.md 레이아웃)
        proposals.json     ({"_meta": {...}, "proposals": [...]})
        review.md          (매핑 경고, 미매칭 팀, 누계 교차검증)

시트의 '현황판' 탭 컬럼과 스킬 입력의 대응은 references/sheet-schema.md 참고.
사람이 개입하지 않는 Routine 에서도 산출물이 나오도록, 확신이 없는 항목은 멈추지 않고
`_unmatched` / `_draft` 플래그와 review.md 경고로 남긴다.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import dumpio  # noqa: E402
import naming  # noqa: E402
import update_status_excel as use  # noqa: E402

SKILL_DIR = Path(__file__).resolve().parent.parent

# 덤프 헤더 라벨(부분일치) → 내부 키
COLUMN_MAP = (
    ("제안월", "제안월"),
    ("본부", "본부"),
    ("사업부", "사업부"),
    ("사업팀", "팀"),
    ("접수자", "접수자"),
    ("소속", "소속"),
    ("제안자", "KOL"),
    ("진료과", "진료과"),
    ("품목 분류", "품목분류"),
    ("개발 형태", "개발형태"),
    ("대분류", "대분류"),
    ("소분류", "소분류"),
    ("주요 적응증", "적응증"),
    ("대표 제품명", "제품명"),
    ("성분", "성분"),
    ("제형", "제형"),
    ("제안 사유", "사유"),
    ("검토 부서", "검토부서"),
    ("담당자", "담당자"),
    ("세부 현황", "세부현황"),
    ("검토 내용", "검토내용"),
)

MONTH_RE = re.compile(r"^\s*\d{2}\s*[년냔]\s*\d{1,2}\s*월\s*$")
NEW_PROPOSAL_HEADERS = ["제안 연월", "본부", "사업부", "팀/사무소", "제안자", "건수",
                        "제안 유형", "아이디어 유형", "제안 내용", "검토 결과", "이메일 표기"]


# ---------- 덤프 파싱 ----------

def build_index(header: list[str]) -> dict[str, int]:
    index: dict[str, int] = {}
    for label, key in COLUMN_MAP:
        for pos, cell in enumerate(header):
            if label in cell and key not in index:
                index[key] = pos
                break
    # 현황(진행중/검토 중단) 컬럼은 헤더가 비어 있어 '세부 현황' 바로 앞 칸으로 잡는다.
    if "세부현황" in index and index["세부현황"] > 0 and not header[index["세부현황"] - 1]:
        index["현황"] = index["세부현황"] - 1
    return index


def _records_from_table(table: list[list[str]], header: list[str], warnings: list[str]) -> list[dict]:
    """헤더 다음 행들(이미 정규화된 셀 리스트)을 내부 키 dict 로."""
    index = build_index(header)
    missing = [k for _, k in COLUMN_MAP if k not in index]
    if missing:
        warnings.append(f"헤더에서 찾지 못한 컬럼: {missing}")
    rows: list[dict] = []
    for cells in table:
        if not any(cells):
            continue
        row = {key: (cells[pos] if pos < len(cells) else "") for key, pos in index.items()}
        if not MONTH_RE.match(row.get("제안월", "")):
            if any(row.values()):
                warnings.append("제안월이 비어 있거나 형식이 달라 건너뜀: " + json.dumps(
                    {k: row[k] for k in ("제안월", "접수자", "제품명") if k in row}, ensure_ascii=False))
            continue
        row["제안월"] = naming.normalize_month(row["제안월"])
        rows.append(row)
    return rows


def _missing_header_exit(anchor_labels):
    raise SystemExit(
        "시트 덤프에서 '현황판' 헤더를 찾지 못했습니다. 시트 구조가 바뀌었는지 확인하세요.\n"
        f"기대한 라벨: {', '.join(anchor_labels)}"
    )


def parse_csv_dump(text: str, anchor_labels) -> tuple[list[dict], list[str]]:
    """download_file_content(text/csv) 결과 — 첫 탭('현황판') 전체."""
    warnings: list[str] = []
    table = dumpio.csv_rows(text)
    hdr = next((i for i, cells in enumerate(table)
                if all(any(label == c for c in cells) for label in anchor_labels)), None)
    if hdr is None:
        _missing_header_exit(anchor_labels)
    rows = _records_from_table(table[hdr + 1:], table[hdr], warnings)
    return rows, warnings


def parse_dump(text: str, anchor_labels) -> tuple[list[dict], list[str]]:
    """덤프(markdown 표 또는 CSV)의 첫 번째 앵커 블록 데이터 행을 dict 목록으로. (rows, warnings)"""
    if dumpio.is_csv_dump(text):
        return parse_csv_dump(text, anchor_labels)
    warnings: list[str] = []
    lines = text.split("\n")
    blocks = dumpio.find_header_blocks(lines, anchor_labels)
    if not blocks:
        _missing_header_exit(anchor_labels)
    header_idx, header = blocks[0]
    table = []
    for line in lines[header_idx + 1:]:
        if "|" not in line:
            break
        cells = dumpio.split_row(line)
        if dumpio.is_separator(cells):
            continue
        table.append([dumpio.unescape(c) for c in cells])
    rows = _records_from_table(table, header, warnings)

    if len(blocks) > 1:
        # 두 번째 탭(구버전 복제)의 행 수를 세어 첫 블록이 더 작으면 경고
        second_idx = blocks[1][0]
        n2 = sum(1 for l in lines[second_idx + 1:] if "|" in l and MONTH_RE.match(
            dumpio.unescape(dumpio.split_row(l)[0]) if dumpio.split_row(l) else ""))
        if n2 > len(rows):
            warnings.append(f"앵커가 {len(blocks)}번 잡혔고 두 번째 블록({n2}행)이 첫 블록({len(rows)}행)보다 큽니다. 탭 순서를 확인하세요.")
    return rows, warnings


# ---------- 정규화 ----------

def norm_bonbu(raw: str, mapping: dict) -> str:
    for rule in mapping.get("bonbu", []):
        if re.search(rule["match"], raw or ""):
            return rule["value"]
    return (raw or "").strip()


def norm_division(raw: str, bonbu: str, mapping: dict) -> str:
    raw = (raw or "").strip()
    suffix = mapping.get("division_suffix", {}).get(bonbu)
    if raw and suffix and not raw.endswith(suffix):
        return raw + suffix
    return raw


def norm_team(raw: str, bonbu: str, mapping: dict) -> str:
    raw = (raw or "").strip()
    if not raw or raw in ("-", "\\-"):
        return ""
    aliases = {use.normalize(k): v for k, v in mapping.get("team_aliases", {}).items()}
    key = use.normalize(raw)
    if key in aliases:
        return aliases[key]
    suffix = mapping.get("team_suffix", {}).get(bonbu)
    if suffix and not raw.endswith(suffix):
        return raw + suffix
    return raw


def norm_review(detail: str, status: str, mapping: dict) -> str:
    rs = mapping.get("review_status", {})
    detail = (detail or "").strip()
    status = (status or "").strip()
    for rule in rs.get("rules", []):
        if rule.get("detail", "") in detail and (not rule.get("status") or rule["status"] == status):
            return rule["value"]
    if status == "검토 중단":
        return rs.get("stopped_value", "검토 중단")
    return detail or status


def _clean(s: str) -> str:
    s = (s or "").strip()
    return "" if s in ("-", "\\-") else s


def draft_content(row: dict, mapping: dict) -> str:
    """제안 내용 초안. Claude 가 세션에서 문구를 다듬는다(사실은 바꾸지 않음)."""
    name = _clean(row.get("제품명"))
    ingr = _clean(row.get("성분"))
    sub = _clean(row.get("소분류"))
    indi = _clean(row.get("적응증"))
    # 성분표 전체가 들어간 경우(경장영양제 등)는 초안에서 잘라낸다 — Claude 가 다듬을 때 _source 참고
    if len(ingr) > 60:
        ingr = ingr[:57].rstrip() + "…"
    if len(indi) > 60:
        indi = indi[:57].rstrip() + "…"
    major = _clean(row.get("대분류"))
    head = name or ingr or indi
    if major in mapping.get("compound_types", []):
        text = f"{head} {major}" if head else major
        if ingr and ingr != head:
            text += f"({ingr})"
    else:
        text = head
        if ingr and ingr != head:
            text += f"({ingr})"
        if sub:
            text += f" {sub}"
        elif major and major not in text:
            text += f" {major}"
        if indi and indi != head:
            text += f"({indi})"
    return re.sub(r"\s{2,}", " ", text).strip()


def to_proposal(row: dict, mapping: dict) -> dict:
    bonbu = norm_bonbu(row.get("본부", ""), mapping)
    team = norm_team(row.get("팀", ""), bonbu, mapping)
    display_suffix = mapping.get("team_display_suffix", {}).get(bonbu, "")
    p = {
        "월": row["제안월"],
        "본부": bonbu,
        "본부표시": mapping.get("bonbu_display", {}).get(bonbu, bonbu),
        "사업부": norm_division(row.get("사업부", ""), bonbu, mapping),
        "팀": team,
        "팀표시": (team + display_suffix) if team and display_suffix and not team.endswith(display_suffix) else team,
        "제안자": _clean(row.get("접수자")),
        "건수": 1,
        "제안유형": f"{_clean(row.get('품목분류'))} {_clean(row.get('개발형태')).replace(' ', '')}".strip(),
        "아이디어유형": _clean(row.get("대분류")),
        "제안내용": draft_content(row, mapping),
        "검토결과": norm_review(row.get("세부현황"), row.get("현황"), mapping),
        "이메일표기": "",
        "_draft": True,
        "_source": {k: row.get(k, "") for k in ("본부", "사업부", "팀", "접수자", "KOL", "소속",
                                                  "품목분류", "개발형태", "대분류", "소분류",
                                                  "적응증", "제품명", "성분", "제형", "현황", "세부현황")},
    }
    return p


# ---------- 베이스 현황판 대조 ----------

def load_base_info(base_path: str | Path) -> dict:
    """베이스 현황판에서 팀명 집합, 시트2 기존 키, 누계 총건수를 읽는다."""
    wb = openpyxl.load_workbook(base_path, data_only=False)
    info = {"team_names": [], "sheet2_keys": set(), "total": None, "prev_text": None}
    ws1 = wb["1. 제안 현황"] if "1. 제안 현황" in wb.sheetnames else None
    if ws1 is not None:
        headers = []
        for r in range(1, ws1.max_row + 1):
            for cell in ws1[r]:
                if cell.value and str(cell.value).strip() in ("팀", "사업부", "사무소"):
                    headers.append((r, cell.column))
                    break
        for i, (hr, col) in enumerate(headers):
            end = headers[i + 1][0] - 1 if i + 1 < len(headers) else ws1.max_row
            for r in range(hr + 1, end + 1):
                v = ws1.cell(row=r, column=col).value
                if v and isinstance(v, str) and not re.match(r"^\d+월$", v.strip()):
                    info["team_names"].append(v.strip())
        for r in range(1, min(ws1.max_row, 40) + 1):
            for c in range(1, ws1.max_column + 1):
                v = ws1.cell(row=r, column=c).value
                if isinstance(v, str) and "누계 진행 결과" in v:
                    info["prev_text"] = v.strip()
                    m = re.search(r"총\s*(\d+)\s*건", v)
                    if m:
                        info["total"] = int(m.group(1))
    sheet2 = next((n for n in wb.sheetnames if n.startswith("2.") and "누적" in n), None)
    if sheet2:
        ws2 = wb[sheet2]
        header_row = None
        for r in range(1, 10):
            if "제안 연월" in [c.value for c in ws2[r]]:
                header_row = r
                break
        if header_row:
            hm = {c.value: c.column for c in ws2[header_row] if c.value}
            for r in range(header_row + 1, ws2.max_row + 1):
                ym = ws2.cell(row=r, column=hm.get("제안 연월", 1)).value
                if not ym:
                    continue
                try:
                    key = (naming.normalize_month(str(ym)),
                           use.normalize(ws2.cell(row=r, column=hm.get("제안자", 3)).value),
                           use.normalize(ws2.cell(row=r, column=hm.get("아이디어 유형", 5)).value))
                except ValueError:
                    continue
                info["sheet2_keys"].add(key)
    return info


def team_matches(team: str, team_names: list[str]) -> bool:
    target = use.normalize(team)
    if not target:
        return False
    for name in team_names:
        n = use.normalize(name)
        if n == target or target in n or n in target:
            return True
    return False


def month_key(label: str) -> tuple[int, int]:
    return naming.parse_month(label)


def dedupe_key(p: dict) -> tuple:
    src = p.get("_source", {})
    return (p["월"], use.normalize(p["제안자"]), use.normalize(p["아이디어유형"]),
            use.normalize(src.get("제품명") or src.get("성분") or p["제안내용"]))


# ---------- 선택 ----------

def select_proposals(rows: list[dict], report_month: str, config: dict,
                     base_info: dict | None, months_override=None) -> tuple[list[dict], list[str]]:
    mapping = config.get("mapping", {})
    sel = config.get("selection", {})
    warnings: list[str] = []
    yy, mm = naming.parse_month(report_month)
    py, pm = (yy, mm - 1) if mm > 1 else (yy - 1, 12)
    prev_label = f"{py}년 {pm}월"
    targets = {naming.normalize_month(m) for m in (months_override or [prev_label])}
    since = month_key(sel.get("catchup_since", "26년 1월"))

    chosen: list[dict] = []
    seen: set = set()
    for row in rows:
        p = to_proposal(row, mapping)
        mk = month_key(p["월"])
        key2 = (p["월"], use.normalize(p["제안자"]), use.normalize(p["아이디어유형"]))
        in_base = base_info is not None and key2 in base_info["sheet2_keys"]
        if p["월"] in targets:
            if in_base:
                warnings.append(f"이미 베이스 현황판 시트2에 있어 건너뜀: {p['월']} {p['팀']} {p['제안자']} {p['아이디어유형']}")
                continue
            p["_reason"] = "전월 신규"
        elif base_info is not None and since <= mk < (py, pm) and not in_base:
            p["_reason"] = "캐치업(베이스 시트2에 없는 과거 행)"
            warnings.append(f"캐치업으로 포함: {p['월']} {p['팀']} {p['제안자']} {p['아이디어유형']} — 이전 달에 누락된 행인지 확인")
        else:
            continue
        dk = dedupe_key(p)
        if dk in seen:
            warnings.append(f"중복 행 제거: {p['월']} {p['제안자']} {p['아이디어유형']} {p['제안내용']}")
            continue
        seen.add(dk)
        if base_info is not None:
            p["_unmatched"] = not team_matches(p["팀"], base_info["team_names"])
            if p["_unmatched"]:
                warnings.append(f"팀 미매칭: '{p['팀']}' (시트 원문 '{row.get('팀','')}', 본부 '{row.get('본부','')}') — config.mapping.team_aliases 확인 필요")
        chosen.append(p)
    return chosen, warnings


# ---------- 출력 ----------

def write_new_proposals_xlsx(proposals: list[dict], path: Path, report_month: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "신규제안"
    ws.cell(row=1, column=1, value=f"BEAR 신규 제안 ({report_month} 보고)")
    for c, h in enumerate(NEW_PROPOSAL_HEADERS, 1):
        ws.cell(row=3, column=c, value=h)
    for r, p in enumerate(proposals, 4):
        vals = [p["월"], p["본부"], p["사업부"], p["팀"], p["제안자"], p["건수"],
                p["제안유형"], p["아이디어유형"], p["제안내용"], p["검토결과"], p.get("이메일표기", "")]
        for c, v in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=v)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def proposals_from_json(path: str | Path) -> tuple[list[dict], dict]:
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if isinstance(data, list):
        return data, {}
    return data.get("proposals", []), data.get("_meta", {})


def write_review(path: Path, meta: dict, proposals: list[dict], warnings: list[str]):
    lines = [f"# BEAR {meta['report_month']} 업데이트 검토 메모", ""]
    lines += [f"- 시트 행 수(전체): {meta['sheet_row_count']}",
              f"- 베이스 현황판: {meta.get('base') or '(없음)'}",
              f"- 베이스 누계 문구: {meta.get('base_prev_text') or '(없음)'}",
              f"- 신규 {meta['new_count']}건 → 총 {meta['total_count']}건 (계산 방식: {meta['total_count_source']})",
              f"- 시트 최종 수정: {meta.get('sheet_modified_time') or '(미기록)'}", ""]
    lines += ["## 신규 제안", ""]
    if not proposals:
        lines.append("(없음)")
    for i, p in enumerate(proposals, 1):
        flag = " **[팀 미매칭]**" if p.get("_unmatched") else ""
        lines.append(f"{i}. {p['월']} {p['본부']} {p['사업부']} **{p['팀']} {p['제안자']}** — {p['아이디어유형']} / {p['제안내용']} / {p['검토결과']} ({p.get('_reason','')}){flag}")
    lines += ["", "## 확인 필요", ""]
    attention = [w for w in warnings if "미매칭" in w or "캐치업" in w or "불일치" in w]
    lines += [f"- {w}" for w in attention] or ["(없음)"]
    lines += ["", "## 기타 경고", ""]
    lines += [f"- {w}" for w in warnings if w not in attention] or ["(없음)"]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


# ---------- main ----------

def run(dump: str, report_month: str, out_dir: str | Path, config: dict,
        base: str | None = None, months=None, sheet_modified_time: str | None = None) -> dict:
    out_dir = Path(out_dir)
    text = dumpio.load_dump(dump)
    rows, warnings = parse_dump(text, config["sheet"]["anchor_labels"])
    base_info = load_base_info(base) if base else None
    if base_info is None:
        warnings.append("베이스 현황판 없음: 팀 매칭·캐치업·누계 교차검증을 건너뜀")
    proposals, w2 = select_proposals(rows, report_month, config, base_info, months)
    warnings += w2

    new_count = sum(int(p.get("건수", 1)) for p in proposals)
    source = config.get("selection", {}).get("total_count_source", "base_plus_new")
    base_total = base_info["total"] if base_info else None
    if source == "sheet_rows" or base_total is None:
        total = len(rows)
        if base_total is None and base_info is not None:
            warnings.append("베이스 시트1에서 '누계 진행 결과 … 총 N건' 문구를 못 찾아 시트 행 수를 누계로 사용")
    else:
        total = base_total + new_count
        if total != len(rows):
            warnings.append(f"누계 불일치: 베이스 {base_total} + 신규 {new_count} = {total} 이지만 시트 행 수는 {len(rows)} — 누계 문구를 확인")

    meta = {
        "report_month": naming.normalize_month(report_month),
        "new_count": new_count,
        "total_count": total,
        "total_count_source": source if base_total is not None else "sheet_rows",
        "sheet_row_count": len(rows),
        "base": str(base) if base else None,
        "base_prev_text": base_info["prev_text"] if base_info else None,
        "sheet_modified_time": sheet_modified_time,
        "warnings": warnings,
        "blocking": False,
    }
    out_dir.mkdir(parents=True, exist_ok=True)
    write_new_proposals_xlsx(proposals, out_dir / "new_proposals.xlsx", meta["report_month"])
    (out_dir / "proposals.json").write_text(
        json.dumps({"_meta": meta, "proposals": proposals}, ensure_ascii=False, indent=2), encoding="utf-8")
    write_review(out_dir / "review.md", meta, proposals, warnings)
    return meta


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dump", required=True, help="read_file_content 결과 파일 (JSON 또는 markdown)")
    ap.add_argument("--report-month", required=True, help="예: '26년 9월'")
    ap.add_argument("--base", help="이전 달 현황판 .xlsx")
    ap.add_argument("--config", default=str(SKILL_DIR / "config.json"))
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--months", nargs="*", help="선택할 제안월 (기본: 보고월의 전월)")
    ap.add_argument("--sheet-modified-time", default=None)
    args = ap.parse_args(argv)
    config = json.loads(Path(args.config).read_text(encoding="utf-8"))
    meta = run(args.dump, args.report_month, args.out_dir, config, args.base, args.months, args.sheet_modified_time)
    print(f"v 신규 {meta['new_count']}건 / 총 {meta['total_count']}건 → {args.out_dir}", file=sys.stderr)
    for w in meta["warnings"]:
        print("[warn] " + w, file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
