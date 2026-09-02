#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""openpyxl 워크시트 → HTML 표 (Chromium 스크린샷용).

LibreOffice Calc 가 없는 샌드박스에서 시트 이미지를 만들기 위한 백엔드.
셀 값·병합·열 너비·행 높이·글꼴(굵기/크기/색)·채우기·테두리·정렬을 CSS 로 옮긴다.

openpyxl 은 수식을 계산하지 않으므로(저장 직후 캐시값 없음) 현황판에 쓰이는 단순 수식
(`=SUM(F14:Q14)`, `=C5+D5` 등)만 여기서 직접 평가한다. 평가 못 하는 수식은 빈칸으로 두고
경고를 남긴다.
"""
from __future__ import annotations

import html
import re
from pathlib import Path

from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

_REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
_RANGE = re.compile(r"\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+")
_SUM = re.compile(r"SUM\(([^()]*)\)", re.I)
_SAFE = re.compile(r"^[\d\.\s\+\-\*/\(\)]+$")


# ---------- 수식 평가 ----------

class FormulaEvaluator:
    def __init__(self, ws):
        self.ws = ws
        self.cache: dict[tuple[int, int], float | None] = {}
        self.warnings: list[str] = []

    def cell_value(self, row: int, col: int):
        key = (row, col)
        if key in self.cache:
            return self.cache[key]
        self.cache[key] = 0  # 순환 참조 방지
        v = self.ws.cell(row=row, column=col).value
        if isinstance(v, str) and v.startswith("="):
            v = self.evaluate(v[1:], f"{get_column_letter(col)}{row}")
        self.cache[key] = v
        return v

    def _num(self, v) -> float:
        if v is None or v == "":
            return 0.0
        if isinstance(v, bool):
            return float(v)
        if isinstance(v, (int, float)):
            return float(v)
        try:
            return float(str(v).replace(",", ""))
        except ValueError:
            return 0.0

    def _sum_args(self, args: str) -> float:
        total = 0.0
        for arg in (a.strip() for a in args.split(",") if a.strip()):
            if _RANGE.fullmatch(arg):
                c1, r1, c2, r2 = range_boundaries(arg.replace("$", ""))
                for r in range(r1, r2 + 1):
                    for c in range(c1, c2 + 1):
                        total += self._num(self.cell_value(r, c))
            elif _REF.fullmatch(arg):
                m = _REF.fullmatch(arg)
                total += self._num(self.cell_value(int(m.group(2)), column_index_from_string(m.group(1))))
            else:
                total += self._num(arg)
        return total

    def evaluate(self, expr: str, where: str = ""):
        try:
            expr = _SUM.sub(lambda m: repr(self._sum_args(m.group(1))), expr)
            expr = _REF.sub(lambda m: repr(self._num(self.cell_value(
                int(m.group(2)), column_index_from_string(m.group(1))))), expr)
            if _SAFE.match(expr):
                val = eval(expr, {"__builtins__": {}}, {})  # 숫자·사칙연산만 남은 문자열
                return int(val) if float(val).is_integer() else val
        except Exception as e:  # noqa: BLE001
            self.warnings.append(f"{where}: 수식 평가 실패 ({e})")
            return None
        self.warnings.append(f"{where}: 지원하지 않는 수식 '={expr}' → 빈칸")
        return None


# ---------- 스타일 ----------

def _rgb(color) -> str | None:
    """openpyxl Color → '#rrggbb' (theme/indexed 색은 None)."""
    if color is None:
        return None
    try:
        if color.type == "rgb" and isinstance(color.rgb, str) and len(color.rgb) >= 6:
            return "#" + color.rgb[-6:]
    except AttributeError:
        return None
    return None


_BORDER_PX = {"thin": "1px", "hair": "1px", "dotted": "1px", "dashed": "1px",
              "medium": "2px", "thick": "3px", "double": "3px"}


def _border_css(border) -> str:
    css = []
    for side in ("top", "right", "bottom", "left"):
        s = getattr(border, side, None)
        if s is not None and s.style:
            px = _BORDER_PX.get(s.style, "1px")
            style = "double" if s.style == "double" else ("dashed" if s.style in ("dashed", "dotted") else "solid")
            css.append(f"border-{side}:{px} {style} {_rgb(s.color) or '#000'}")
    return ";".join(css)


def _fmt_value(v, number_format: str) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        if "%" in (number_format or ""):
            return f"{v * 100:.0f}%"
        if v.is_integer():
            return str(int(v))
        return f"{v:.2f}".rstrip("0").rstrip(".")
    if isinstance(v, int):
        if "%" in (number_format or ""):
            return f"{v * 100:.0f}%"
        return str(v)
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    return str(v)


# ---------- HTML ----------

def sheet_to_html(ws, font_dir: Path, stop_regex: str | None = None,
                  max_rows: int | None = None, base_font_px: int = 13) -> tuple[str, list[str]]:
    """워크시트를 HTML 문서 문자열로. (html, warnings)"""
    ev = FormulaEvaluator(ws)
    warnings: list[str] = []

    # 렌더 범위: 값이 있는 마지막 행/열까지. stop_regex 에 맞는 행 앞에서 끊는다.
    max_r, max_c = ws.max_row, ws.max_column
    last_r = 0
    last_c = 0
    for row in ws.iter_rows(min_row=1, max_row=max_r, max_col=max_c):
        for cell in row:
            if cell.value not in (None, ""):
                last_r = max(last_r, cell.row)
                last_c = max(last_c, cell.column)
    if stop_regex:
        pat = re.compile(stop_regex)
        for r in range(1, last_r + 1):
            if any(isinstance(c.value, str) and pat.search(c.value.strip()) for c in ws[r]):
                last_r = r - 1
                break
        while last_r > 0 and all(ws.cell(row=last_r, column=c).value in (None, "") for c in range(1, last_c + 1)):
            last_r -= 1
    if max_rows:
        last_r = min(last_r, max_rows)
    if last_r == 0 or last_c == 0:
        return "<html><body></body></html>", ["빈 시트"]

    # 병합 셀
    span: dict[tuple[int, int], tuple[int, int]] = {}
    covered: set[tuple[int, int]] = set()
    for rng in ws.merged_cells.ranges:
        c1, r1, c2, r2 = rng.bounds
        if r1 > last_r:
            continue
        span[(r1, c1)] = (min(r2, last_r) - r1 + 1, min(c2, last_c) - c1 + 1)
        for r in range(r1, min(r2, last_r) + 1):
            for c in range(c1, min(c2, last_c) + 1):
                if (r, c) != (r1, c1):
                    covered.add((r, c))

    # 열 너비(px). 열 너비가 지정되지 않은 시트(테스트용 합성 파일 등)에서 긴 텍스트가 잘리지 않도록
    # 내용 길이에 맞춰 autofit 하되 상한을 둔다(넘치면 wrap).
    widths = []
    need = {}
    for row in ws.iter_rows(min_row=1, max_row=last_r, max_col=last_c):
        for cell in row:
            if cell.value in (None, "") or (cell.row, cell.column) in span and span[(cell.row, cell.column)][1] > 1:
                continue
            s = str(cell.value)
            if s.startswith("="):
                s = "0000"
            units = sum(2 if ord(ch) > 0x2E7F else 1 for ch in s)
            need[cell.column] = max(need.get(cell.column, 0), units)
    force_wrap: set[int] = set()
    for c in range(1, last_c + 1):
        dim = ws.column_dimensions.get(get_column_letter(c))
        if dim is not None and dim.hidden:
            widths.append(0)
            continue
        explicit = dim is not None and dim.width
        w = dim.width if explicit else 8.43
        px = int(w * 7 + 5)
        if not explicit and c in need:
            want = int(need[c] * 7 + 12)
            if want > px:
                px = min(want, 480)
                if want > 480:
                    force_wrap.add(c)
        widths.append(px)

    font_dir = Path(font_dir)
    css = f"""
@font-face{{font-family:'NanumGothic';src:url('file://{font_dir / 'NanumGothic.ttf'}');font-weight:normal}}
@font-face{{font-family:'NanumGothic';src:url('file://{font_dir / 'NanumGothicBold.ttf'}');font-weight:bold}}
body{{margin:0;padding:12px;background:#fff;font-family:'NanumGothic','Malgun Gothic',sans-serif;font-size:{base_font_px}px;color:#000}}
table{{border-collapse:collapse;table-layout:fixed}}
td{{padding:1px 4px;overflow:hidden;vertical-align:middle;white-space:nowrap;line-height:1.35}}
td.wrap{{white-space:normal;word-break:break-all}}
td.spill{{overflow:visible;position:relative;z-index:1}}
"""
    total_w = sum(widths) + 8 * len(widths) + 40
    parts = [f"<!doctype html><html><head><meta charset='utf-8'><style>{css}</style></head>"
             f"<body data-width='{total_w}'><table style='width:{sum(widths) + 8 * len(widths)}px'>"]
    parts.append("<colgroup>" + "".join(f"<col style='width:{w}px'>" for w in widths) + "</colgroup>")
    for r in range(1, last_r + 1):
        rd = ws.row_dimensions.get(r)
        if rd is not None and rd.hidden:
            continue
        h_pt = rd.height if rd is not None and rd.height else 15
        parts.append(f"<tr style='height:{int(h_pt * 96 / 72)}px'>")
        for c in range(1, last_c + 1):
            if (r, c) in covered or widths[c - 1] == 0:
                continue
            cell = ws.cell(row=r, column=c)
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                v = ev.evaluate(v[1:], cell.coordinate)
            text = html.escape(_fmt_value(v, cell.number_format))
            rs, cs = span.get((r, c), (1, 1))
            styles = []
            f = cell.font
            if f is not None:
                if f.bold:
                    styles.append("font-weight:bold")
                if f.italic:
                    styles.append("font-style:italic")
                if f.size and abs(float(f.size) - 11) > 0.5:
                    styles.append(f"font-size:{float(f.size) * base_font_px / 11:.1f}px")
                col = _rgb(f.color)
                if col and col != "#000000":
                    styles.append(f"color:{col}")
                if f.underline:
                    styles.append("text-decoration:underline")
            fill = cell.fill
            if fill is not None and fill.fill_type == "solid":
                bg = _rgb(fill.fgColor)
                if bg and bg != "#000000":
                    styles.append(f"background:{bg}")
            al = cell.alignment
            cls = ""
            if al is not None:
                if al.horizontal in ("center", "centerContinuous"):
                    styles.append("text-align:center")
                elif al.horizontal == "right":
                    styles.append("text-align:right")
                elif al.horizontal == "left":
                    styles.append("text-align:left")
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    styles.append("text-align:right")
                if al.vertical == "top":
                    styles.append("vertical-align:top")
                elif al.vertical == "bottom":
                    styles.append("vertical-align:bottom")
                if al.wrap_text:
                    cls = " class='wrap'"
            elif isinstance(v, (int, float)) and not isinstance(v, bool):
                styles.append("text-align:right")
            if c in force_wrap and not cls:
                cls = " class='wrap'"
            # Excel 처럼: 오른쪽 이웃 셀이 비어 있으면 텍스트가 그쪽으로 넘쳐 보이게 (제목·섹션 표기 잘림 방지)
            if (not cls and isinstance(v, str) and rs == 1 and cs == 1
                    and c < last_c and ws.cell(row=r, column=c + 1).value in (None, "")
                    and (al is None or al.horizontal in (None, "general", "left"))):
                cls = " class='spill'"
            b = _border_css(cell.border)
            if b:
                styles.append(b)
            attrs = (f" rowspan='{rs}'" if rs > 1 else "") + (f" colspan='{cs}'" if cs > 1 else "")
            parts.append(f"<td{attrs}{cls} style='{';'.join(styles)}'>{text}</td>")
        parts.append("</tr>")
    parts.append("</table></body></html>")
    warnings += ev.warnings
    return "".join(parts), warnings


def estimate_size(ws, last_r: int | None = None, html_doc: str | None = None) -> tuple[int, int]:
    """스크린샷 창 크기 추정 (가로 px, 세로 px). html_doc 이 있으면 그 안의 data-width 를 쓴다."""
    width = 40
    m = re.search(r"data-width='(\d+)'", html_doc or "")
    if m:
        width = int(m.group(1))
    else:
        for c in range(1, ws.max_column + 1):
            dim = ws.column_dimensions.get(get_column_letter(c))
            if dim is not None and dim.hidden:
                continue
            w = dim.width if dim is not None and dim.width else 8.43
            width += int(w * 7 + 13)
    height = 60
    for r in range(1, (last_r or ws.max_row) + 1):
        rd = ws.row_dimensions.get(r)
        h_pt = rd.height if rd is not None and rd.height else 15
        height += int(h_pt * 96 / 72)
    return min(max(width, 600), 4000), min(max(int(height * 1.8), 400), 16000)
